import io

import pytest
from docx import Document

from corretor import ErroCorrecao
from docx_io import extrair_texto, gerar


def _docx(*paragrafos) -> bytes:
    """Cada parágrafo é uma lista de (texto, itálico)."""
    documento = Document()
    for runs in paragrafos:
        p = documento.add_paragraph()
        for texto, italico in runs:
            p.add_run(texto).italic = italico
    saida = io.BytesIO()
    documento.save(saida)
    return saida.getvalue()


def _runs(dados: bytes) -> list[list[tuple[str, bool]]]:
    return [[(r.text, bool(r.italic)) for r in p.runs] for p in Document(io.BytesIO(dados)).paragraphs]


def test_extrai_um_paragrafo_por_bloco():
    dados = _docx([("Era uma ", False), ("vez", True)], [], [("Fim.", False)])
    assert extrair_texto(dados) == "Era uma vez\n\n\n\nFim."


def test_devolve_correcoes_mantendo_a_formatacao():
    dados = _docx([("Ele tava ", False), ("cansadu", True), (" demais.", False)], [("Fim.", False)])
    novo, preservada = gerar("Ele estava cansado demais.\n\nFim.", dados)
    assert preservada
    assert _runs(novo) == [[("Ele estava ", False), ("cansado", True), (" demais.", False)], [("Fim.", False)]]


def test_troca_de_palavra_inteira_em_italico_continua_em_italico():
    dados = _docx([("Ela viu ", False), ("noslençois", True), (".", False)])
    novo, _ = gerar("Ela viu nos lençóis.", dados)
    assert _runs(novo) == [[("Ela viu ", False), ("nos lençóis", True), (".", False)]]


def test_quebra_de_linha_manual_e_mantida():
    documento = Document()
    documento.add_paragraph("").add_run("Linha um\nlinha doiz")
    saida = io.BytesIO()
    documento.save(saida)
    texto = extrair_texto(saida.getvalue())
    assert texto == "Linha um linha doiz"
    novo, _ = gerar("Linha um linha dois", saida.getvalue())
    assert Document(io.BytesIO(novo)).paragraphs[0].runs[0].text == "Linha um\nlinha dois"


def test_paragrafos_diferentes_geram_documento_simples():
    dados = _docx([("Um.", False)])
    novo, preservada = gerar("Um.\n\nDois.", dados)
    assert not preservada
    assert [p.text for p in Document(io.BytesIO(novo)).paragraphs] == ["Um.", "Dois."]


def test_sem_original_gera_um_paragrafo_por_linha():
    novo, preservada = gerar("Um.\nDois.\n\nTrês.")
    assert not preservada
    assert [p.text for p in Document(io.BytesIO(novo)).paragraphs] == ["Um.", "Dois.", "Três."]


def test_arquivo_invalido_da_erro_amigavel():
    with pytest.raises(ErroCorrecao, match="Word"):
        extrair_texto(b"isto nao e um docx")
