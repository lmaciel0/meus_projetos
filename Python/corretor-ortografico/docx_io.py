"""Entrada e saída em Word (.docx), preservando a formatação dos parágrafos."""

import difflib
import io
import re
from zipfile import BadZipFile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from corretor import ErroCorrecao

# Cada parágrafo do Word vira um parágrafo separado por linha em branco, como no PDF.
SEPARADOR = "\n\n"


def _texto(paragrafo) -> str:
    # Quebra de linha manual (Shift+Enter) vira espaço: tem o mesmo tamanho, então as
    # posições continuam valendo para devolver as correções aos trechos certos.
    return "".join(r.text for r in paragrafo.runs).replace("\n", " ")


def _abrir(dados: bytes):
    try:
        return Document(io.BytesIO(dados))
    except (PackageNotFoundError, BadZipFile, KeyError, ValueError, OSError) as e:
        raise ErroCorrecao(f"Não foi possível ler o arquivo do Word: {e}") from e


def extrair_texto(dados: bytes) -> str:
    """Texto dos parágrafos do corpo do documento (tabelas, cabeçalhos e notas ficam de fora)."""
    texto = SEPARADOR.join(_texto(p) for p in _abrir(dados).paragraphs)
    if not texto.strip():
        raise ErroCorrecao("Não foi encontrado texto no documento do Word.")
    return texto


def _trocar(textos: list[str], i1: int, i2: int, inserir: str) -> None:
    """Troca o trecho [i1, i2) do parágrafo por `inserir`, mexendo só nos runs afetados."""
    inicios, posicao = [], 0
    for t in textos:
        inicios.append(posicao)
        posicao += len(t)
    faixas = [(k, inicios[k], inicios[k] + len(t)) for k, t in enumerate(textos) if t]
    if i2 > i1:  # troca: o texto novo herda a formatação do primeiro caractere trocado
        alvo = next(k for k, ini, fim in faixas if ini <= i1 < fim)
    elif i1 > 0:  # inserção: continua a formatação do que vem antes ("palavra" + "s")
        alvo = next(k for k, ini, fim in faixas if ini < i1 <= fim)
    else:
        alvo = faixas[0][0] if faixas else 0
    local = i1 - inicios[alvo]
    for k, ini, fim in faixas:
        a, b = max(i1, ini), min(i2, fim)
        if a < b:
            textos[k] = textos[k][: a - ini] + textos[k][b - ini :]
    textos[alvo] = textos[alvo][:local] + inserir + textos[alvo][local:]


def _substituir(paragrafo, novo: str) -> None:
    runs = paragrafo.runs
    antigo = "".join(r.text for r in runs).replace("\n", " ")
    if antigo == novo:
        return
    if not runs:
        paragrafo.add_run(novo)
        return
    textos = [r.text for r in runs]
    operacoes = difflib.SequenceMatcher(None, antigo, novo, autojunk=False).get_opcodes()
    # De trás para frente: as posições anteriores continuam valendo.
    for tag, i1, i2, j1, j2 in reversed(operacoes):
        if tag != "equal":
            _trocar(textos, i1, i2, novo[j1:j2])
    for r, t in zip(runs, textos):
        if r.text != t:
            r.text = t


def _bytes(documento) -> bytes:
    saida = io.BytesIO()
    documento.save(saida)
    return saida.getvalue()


def gerar(texto: str, original: bytes | None = None) -> tuple[bytes, bool]:
    """Gera o .docx do texto corrigido. Retorna (dados, formatação preservada).

    Com o .docx original, as correções são aplicadas nele, mantendo estilos, itálicos e
    tudo mais. Se o número de parágrafos mudou (edição manual), ou se não há original,
    gera um documento simples, um parágrafo por linha.
    """
    if original:
        documento = _abrir(original)
        novos = texto.split(SEPARADOR)
        if len(novos) == len(documento.paragraphs):
            for paragrafo, novo in zip(documento.paragraphs, novos):
                _substituir(paragrafo, novo)
            return _bytes(documento), True

    documento = Document()
    for linha in re.split(r"\n(?:[ \t]*\n)*", texto):
        documento.add_paragraph(linha)
    return _bytes(documento), False
